"""Écriture des résultats dans l'onglet cible du fichier Excel partagé sur Google Drive."""

import os
from datetime import datetime
from io import BytesIO
from typing import Optional

import streamlit as st
from google.auth.transport.requests import AuthorizedSession
from google.oauth2.service_account import Credentials
from openpyxl import load_workbook

_SCOPES = ["https://www.googleapis.com/auth/drive"]
_DRIVE_FILES_URL = "https://www.googleapis.com/drive/v3/files"
_UPLOAD_URL = "https://www.googleapis.com/upload/drive/v3/files"
_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_SHEET_NAME = "cartographie des metiers"

_session: Optional[AuthorizedSession] = None


def _load_credentials() -> Credentials:
    if "gcp_service_account" in st.secrets:
        return Credentials.from_service_account_info(
            dict(st.secrets["gcp_service_account"]), scopes=_SCOPES
        )
    service_account_path = os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"]
    return Credentials.from_service_account_file(service_account_path, scopes=_SCOPES)


def _get_session() -> AuthorizedSession:
    global _session
    if _session is None:
        _session = AuthorizedSession(_load_credentials())
    return _session


def _find_worksheet(workbook, name: str):
    normalized = name.strip().casefold()
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().casefold() == normalized:
            return workbook[sheet_name]
    available = ", ".join(workbook.sheetnames)
    raise KeyError(f"Aucun onglet nommé « {name} » trouvé. Onglets disponibles : {available}")


def append_row(fields: dict) -> None:
    """Télécharge le fichier Excel partagé, ajoute une ligne dans l'onglet cible, puis le ré-uploade."""
    file_id = st.secrets.get("GOOGLE_SHEET_ID") or os.environ["GOOGLE_SHEET_ID"]
    session = _get_session()

    download = session.get(f"{_DRIVE_FILES_URL}/{file_id}?alt=media")
    download.raise_for_status()

    workbook = load_workbook(BytesIO(download.content))
    worksheet = _find_worksheet(workbook, _SHEET_NAME)

    worksheet.append(
        [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            fields["nom"],
            fields["prenom"],
            fields["formation"],
            fields["projet_1"],
            fields["projet_2"],
            fields["alignement_projet"],
            fields["emploi_formation"],
            fields["secteur"],
        ]
    )

    buffer = BytesIO()
    workbook.save(buffer)

    upload = session.patch(
        f"{_UPLOAD_URL}/{file_id}?uploadType=media",
        data=buffer.getvalue(),
        headers={"Content-Type": _MIME_XLSX},
    )
    upload.raise_for_status()
